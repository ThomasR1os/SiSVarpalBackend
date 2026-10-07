from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from catalogos.models import Cliente, PlantillaImportacion, Punto, Vehiculo
from operaciones.models import Destino, Ruta
from usuarios.models import Rol, Usuario

from .destinos import crear_destino
from .rutas import crear_ruta

ALIAS_TIPO = {
    "TRASLADOS": Destino.TipoServicio.TRASLADO,
    "TRASLADO": Destino.TipoServicio.TRASLADO,
    "ENTREGA": Destino.TipoServicio.ENTREGA,
    "ENTREGAS": Destino.TipoServicio.ENTREGA,
    "INTERCAMBIO": Destino.TipoServicio.INTERCAMBIO,
    "RECOJO": Destino.TipoServicio.RECOJO,
}

CAMPOS_DESTINO = {
    "codigo_externo",
    "fecha",
    "tipo_servicio",
    "motivo_servicio",
    "documento_receptor",
    "nombre_receptor",
    "apellido_receptor",
    "telefono_receptor",
    "direccion",
    "distrito",
    "referencia",
    "latitud",
    "longitud",
    "observaciones",
    "codigo_sede",
}


def leer_hoja(archivo):
    try:
        libro = load_workbook(archivo, data_only=True)
    except Exception as exc:
        raise ValidationError({"archivo": ["No se pudo leer el Excel."]}) from exc
    hoja = libro.active
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        raise ValidationError({"archivo": ["El archivo no tiene filas."]})
    encabezados = ["" if celda is None else str(celda).strip() for celda in filas[0]]
    if not any(encabezados):
        raise ValidationError({"archivo": ["El archivo no tiene encabezados."]})
    return encabezados, filas[1:]


def mapeo_de(cliente, tipo):
    plantilla = (
        PlantillaImportacion.objects.filter(cliente=cliente, tipo=tipo, activa=True).order_by("-id").first()
    )
    if plantilla is None:
        return {}
    return {str(origen).strip(): str(destino).strip() for origen, destino in plantilla.mapeo_columnas.items()}


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValueError("Fecha inválida. Use AAAA-MM-DD.")


def _numero(valor):
    if valor is None or _texto(valor) == "":
        return None
    texto = str(valor).strip().replace(",", ".")
    try:
        Decimal(texto)
    except InvalidOperation as exc:
        raise ValueError("Coordenada inválida.") from exc
    return texto


def _fila_vacia(valores):
    return all(valor is None or _texto(valor) == "" for valor in valores)


def _columnas(encabezados, valores, mapeo):
    datos = {}
    extra = {}
    for encabezado, valor in zip(encabezados, valores):
        if not encabezado:
            continue
        if mapeo:
            if encabezado not in mapeo:
                if _texto(valor):
                    extra[encabezado] = _texto(valor)
                continue
            canonico = mapeo[encabezado]
        else:
            canonico = encabezado
        datos[canonico] = valor
    return datos, extra


def _errores_de(exc):
    detalle = getattr(exc, "detail", None)
    if isinstance(detalle, dict):
        mensajes = []
        for clave, valor in detalle.items():
            if isinstance(valor, list):
                mensajes.extend(f"{clave}: {item}" for item in valor)
            else:
                mensajes.append(f"{clave}: {valor}")
        return mensajes or ["Fila inválida."]
    if isinstance(detalle, list):
        return [str(item) for item in detalle]
    return [str(exc)]


def importar_destinos(*, cliente, archivo, usuario):
    encabezados, filas = leer_hoja(archivo)
    mapeo = mapeo_de(cliente, PlantillaImportacion.Tipo.DESTINOS)
    creados = []
    rechazados = []
    vistos = set()
    for numero, valores in enumerate(filas, start=2):
        if _fila_vacia(valores):
            continue
        try:
            datos, extra = _columnas(encabezados, valores, mapeo)
            codigo = _texto(datos.get("codigo_externo"))
            if not codigo:
                raise ValidationError({"codigo_externo": ["El código es obligatorio."]})
            if codigo in vistos:
                raise ValidationError({"codigo_externo": ["El código está repetido en el archivo."]})
            vistos.add(codigo)
            tipo = ALIAS_TIPO.get(_texto(datos.get("tipo_servicio")).upper())
            if tipo is None:
                raise ValidationError({"tipo_servicio": ["Tipo de servicio no reconocido."]})
            payload = {
                "codigo_externo": codigo,
                "fecha": _fecha(datos.get("fecha")),
                "tipo_servicio": tipo,
                "motivo_servicio": _texto(datos.get("motivo_servicio")),
                "documento_receptor": _texto(datos.get("documento_receptor")),
                "nombre_receptor": _texto(datos.get("nombre_receptor")),
                "apellido_receptor": _texto(datos.get("apellido_receptor")),
                "telefono_receptor": _texto(datos.get("telefono_receptor")),
                "direccion": _texto(datos.get("direccion")),
                "distrito": _texto(datos.get("distrito")),
                "referencia": _texto(datos.get("referencia")),
                "observaciones": _texto(datos.get("observaciones")),
                "latitud": _numero(datos.get("latitud")),
                "longitud": _numero(datos.get("longitud")),
                "datos_extra": extra,
                "sede": None,
            }
            if tipo == Destino.TipoServicio.TRASLADO:
                codigo_sede = _texto(datos.get("codigo_sede"))
                sede = Punto.objects.filter(cliente=cliente, codigo=codigo_sede, es_sede=True, activo=True).first()
                if sede is None:
                    raise ValidationError({"codigo_sede": ["No existe una sede activa con ese código."]})
                payload["sede"] = sede
            with transaction.atomic():
                destino = crear_destino(cliente=cliente, datos=payload, usuario=usuario)
            creados.append(destino.id)
        except (ValidationError, ValueError) as exc:
            rechazados.append({"fila": numero, "errores": _errores_de(exc)})
    return {"creados": len(creados), "destinos": creados, "rechazados": rechazados}


def _documentos(valor):
    texto = _texto(valor)
    if not texto:
        return []
    return [parte.strip() for parte in texto.replace(";", ",").split(",") if parte.strip()]


def importar_rutas(*, archivo, usuario):
    encabezados, filas = leer_hoja(archivo)
    creados = []
    rechazados = []
    grupos = {}
    for numero, valores in enumerate(filas, start=2):
        if _fila_vacia(valores):
            continue
        datos, _extra = _columnas(encabezados, valores, {})
        codigo_ruta = _texto(datos.get("codigo_ruta"))
        if not codigo_ruta:
            rechazados.append({"fila": numero, "errores": ["codigo_ruta: El código de ruta es obligatorio."]})
            continue
        grupos.setdefault(codigo_ruta, []).append((numero, datos))

    for codigo_ruta, entradas in grupos.items():
        try:
            with transaction.atomic():
                ruta = _crear_grupo(codigo_ruta, entradas, usuario)
            creados.append(ruta.id)
        except (ValidationError, ValueError) as exc:
            mensajes = _errores_de(exc)
            for numero, _datos in entradas:
                rechazados.append({"fila": numero, "errores": mensajes})
    return {"creados": len(creados), "rutas": creados, "rechazados": rechazados}


def _crear_grupo(codigo_ruta, entradas, usuario):
    primera_fila, primera = entradas[0]
    ruc = _texto(primera.get("ruc"))
    cliente = Cliente.objects.filter(ruc=ruc, activo=True).first()
    if cliente is None:
        raise ValidationError({"ruc": [f"No existe un cliente activo con el RUC {ruc}."]})
    fecha = _fecha(primera.get("fecha"))
    placa = _texto(primera.get("placa")).upper().replace(" ", "").replace("-", "")
    vehiculo = Vehiculo.objects.filter(placa=placa, activo=True).first()
    if vehiculo is None:
        raise ValidationError({"placa": [f"No existe un vehículo activo con la placa {placa}."]})
    conductor = Usuario.objects.filter(
        documento=_texto(primera.get("documento_conductor")),
        rol__nombre=Rol.Nombre.CONDUCTOR,
        is_active=True,
    ).first()
    if conductor is None:
        raise ValidationError({"documento_conductor": ["No existe un conductor activo con ese documento."]})
    documentos = []
    for _numero, datos in entradas:
        documentos.extend(_documentos(datos.get("documentos_auxiliares")))
    auxiliares = []
    vistos = set()
    for documento in documentos:
        if documento in vistos:
            continue
        vistos.add(documento)
        auxiliar = Usuario.objects.filter(documento=documento, rol__nombre=Rol.Nombre.AUXILIAR, is_active=True).first()
        if auxiliar is None:
            raise ValidationError({"documentos_auxiliares": [f"No existe el auxiliar {documento}."]})
        auxiliares.append(auxiliar)
    codigo_origen = _texto(primera.get("codigo_base_origen"))
    codigo_final = _texto(primera.get("codigo_base_final"))
    base_origen = None
    base_final = None
    if codigo_origen:
        base_origen = Punto.objects.filter(cliente=cliente, codigo=codigo_origen, es_base_origen=True, activo=True).first()
        if base_origen is None:
            raise ValidationError({"codigo_base_origen": ["No existe esa base de origen."]})
    if codigo_final:
        base_final = Punto.objects.filter(cliente__isnull=True, codigo=codigo_final, es_base_final=True, activo=True).first()
        if base_final is None:
            raise ValidationError({"codigo_base_final": ["No existe esa base final de la empresa."]})
    for _numero, datos in entradas[1:]:
        if _texto(datos.get("ruc")) not in {"", ruc}:
            raise ValidationError({"codigo_ruta": [f"La ruta {codigo_ruta} mezcla clientes."]})
        fecha_fila = datos.get("fecha")
        if _texto(fecha_fila) and _fecha(fecha_fila) != fecha:
            raise ValidationError({"codigo_ruta": [f"La ruta {codigo_ruta} mezcla fechas."]})
        if _texto(datos.get("placa")).upper().replace(" ", "").replace("-", "") not in {"", placa}:
            raise ValidationError({"placa": [f"La ruta {codigo_ruta} mezcla placas."]})
    pares = []
    for numero, datos in entradas:
        codigo_destino = _texto(datos.get("codigo_externo"))
        destino = Destino.objects.filter(cliente=cliente, codigo_externo=codigo_destino, fecha=fecha).first()
        if destino is None:
            raise ValidationError({"codigo_externo": [f"Fila {numero}: no existe el destino {codigo_destino} en esa fecha."]})
        try:
            orden = int(float(datos.get("orden")))
        except (TypeError, ValueError) as exc:
            raise ValidationError({"orden": [f"Fila {numero}: el orden es obligatorio."]}) from exc
        pares.append((orden, destino))
    pares.sort(key=lambda item: item[0])
    if len({orden for orden, _destino in pares}) != len(pares):
        raise ValidationError({"orden": ["El orden de la ruta está repetido."]})
    destinos = [destino for _orden, destino in pares]
    ruta = crear_ruta(
        cliente=cliente,
        fecha=fecha,
        conductor=conductor,
        vehiculo=vehiculo,
        auxiliares=auxiliares,
        base_origen=base_origen,
        base_final=base_final,
        destinos=destinos,
        creado_por=usuario,
        optimizar=False,
    )
    return ruta
